import argparse
import os
import sys

SKILL_SCRIPTS = os.path.expanduser('~/.config/chipagents/config2/skills/pptx-generator/scripts')
sys.path.insert(0, SKILL_SCRIPTS)

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from pptx_template import (TEMPLATE_PATH, convert_potx_to_pptx, clear_slides, get_placeholder,
                           remove_unused_placeholders, apply_footer, set_placeholder_text, verify_pptx)

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
RED = RGBColor(0xB3, 0x26, 0x1E)
AMBER = RGBColor(0xB2, 0x6A, 0x00)
GRAY = RGBColor(0x60, 0x60, 0x60)
CODE_BG = RGBColor(0x1E, 0x1E, 0x1E)
CODE_FG = RGBColor(0xD4, 0xD4, 0xD4)
DIFF_ADD = RGBColor(0x6A, 0xD3, 0x6F)
DIFF_DEL = RGBColor(0xF2, 0x6D, 0x6D)
MONO = 'Consolas'
A1, A2, A3, A4 = (MSO_THEME_COLOR.ACCENT_1, MSO_THEME_COLOR.ACCENT_2,
                  MSO_THEME_COLOR.ACCENT_3, MSO_THEME_COLOR.ACCENT_4)
D1, D2 = MSO_THEME_COLOR.DARK_1, MSO_THEME_COLOR.DARK_2
BOXES = []


def track(name, x, y, w, h):
    BOXES.append((name, x, y, w, h))


def check_layout(slide_name):
    for i, (n1, x1, y1, w1, h1) in enumerate(BOXES):
        if x1 < 0.3 or y1 < 1.3 or x1 + w1 > 13.05 or y1 + h1 > 6.95:
            print(f'  WARN [{slide_name}] {n1} outside safe area')
        for n2, x2, y2, w2, h2 in BOXES[i + 1:]:
            if not (x1 + w1 <= x2 or x2 + w2 <= x1 or y1 + h1 <= y2 or y2 + h2 <= y1):
                print(f'  WARN [{slide_name}] {n1} overlaps {n2}')
    BOXES.clear()


def _color(font, color):
    if isinstance(color, RGBColor):
        font.color.rgb = color
    else:
        font.color.theme_color = color


def _style_text(tf, lines, size, bold, color, align, font=None, body_size=None):
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.font.size = Pt(size if i == 0 or body_size is None else body_size)
        p.font.bold = bold if i == 0 else False
        if font:
            p.font.name = font
        _color(p.font, color)


def _fill(shape, fill, brightness=0.0):
    shape.fill.solid()
    if isinstance(fill, RGBColor):
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.fore_color.theme_color = fill
        shape.fill.fore_color.brightness = brightness


def box(slide, name, x, y, w, h, lines, fill, brightness=0.0, size=13, bold=True, text_color=WHITE,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, font=None, align=PP_ALIGN.CENTER, line_color=None,
        tracked=True, body_size=None, anchor=MSO_ANCHOR.MIDDLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    _fill(s, fill, brightness)
    if line_color is None:
        s.line.fill.background()
    else:
        _color(s.line, line_color)
        s.line.width = Pt(1.5)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    tf.vertical_anchor = anchor
    _style_text(tf, lines if isinstance(lines, list) else [lines], size, bold, text_color, align, font,
                body_size)
    if tracked:
        track(name, x, y, w, h)
    return s


def container(slide, x, y, w, h, label, color):
    lbl = slide.shapes.add_textbox(Inches(x), Inches(y - 0.36), Inches(w), Inches(0.32))
    _style_text(lbl.text_frame, [label], 13, True, color, PP_ALIGN.LEFT)
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.04
    _fill(s, color, 0.85)
    _color(s.line, color)
    s.line.width = Pt(1.25)
    return s


def text(slide, x, y, w, h, lines, size=11, bold=False, color=D1, align=PP_ALIGN.LEFT, font=None,
         name=None, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.text_frame.margin_left = tb.text_frame.margin_right = Inches(0.02)
    _style_text(tb.text_frame, lines if isinstance(lines, list) else [lines], size, bold, color, align, font)
    if italic:
        for p in tb.text_frame.paragraphs:
            p.font.italic = True
    if name:
        track(name, x, y, w, h)
    return tb


def code(slide, name, x, y, w, h, rows, size=11, title=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    _fill(s, CODE_BG)
    s.line.fill.background()
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.12)
    allrows = ([(title, WHITE)] if title else []) + rows
    for i, (t, c) in enumerate(allrows):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = t
        p.font.name = MONO
        p.font.size = Pt(size)
        p.font.color.rgb = c
        p.font.bold = bool(title) and i == 0
        p.alignment = PP_ALIGN.LEFT
    track(name, x, y, w, h)
    return s


def arrow(slide, x, y, w, h, direction='right', color=D2):
    shp = {'right': MSO_SHAPE.RIGHT_ARROW, 'left': MSO_SHAPE.LEFT_ARROW,
           'down': MSO_SHAPE.DOWN_ARROW, 'up': MSO_SHAPE.UP_ARROW}[direction]
    a = slide.shapes.add_shape(shp, Inches(x), Inches(y), Inches(w), Inches(h))
    _fill(a, color)
    a.line.fill.background()
    return a


def badge(slide, x, y, n, color=A1, d=0.42):
    b = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    _fill(b, color)
    b.line.color.rgb = WHITE
    b.line.width = Pt(1.5)
    tf = b.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _style_text(tf, [str(n)], 13, True, WHITE, PP_ALIGN.CENTER)
    return b


def content_slide(prs, title, layout=2):
    slide = prs.slides.add_slide(prs.slide_layouts[layout])
    set_placeholder_text(slide, 0, title, font_size=26, bold=True)
    for idx in (10, 11):
        ph = get_placeholder(slide, idx)
        if ph is not None:
            ph._element.getparent().remove(ph._element)
    return slide


def finish(prs, slide, name):
    apply_footer(slide, prs)
    remove_unused_placeholders(slide)
    check_layout(name)


def source(slide, txt):
    text(slide, 0.55, 6.62, 12.25, 0.3, txt, 9, color=GRAY, italic=True, name='source')


def slide_cover(prs):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    set_placeholder_text(s, 10, 'Why I Use ChipAgents Web')
    set_placeholder_text(s, 14, 'The ChipAgents agent in a browser UI')
    set_placeholder_text(s, 15, 'ChipAgents v0.44.3 · WSL2 on Windows')
    finish(prs, s, 'cover')


def slide_intro_ca(prs):
    s = content_slide(prs, 'Intro: what is ChipAgents?')
    box(s, 'quote', 0.55, 1.6, 12.25, 0.95, ['“A fully autonomous terminal-based agent for hardware design, '
        'verification, and debugging tasks.”'], A1, 0.85, 18, bold=False, text_color=A1, line_color=A1)
    cards = [('Reads specs and code', 'Searches the repository and manages its own context'),
             ('Executes tasks', 'Works from natural-language instructions with minimal supervision'),
             ('Customizable', 'Rules, coding style, internal build tools and libraries')]
    for i, (h, b) in enumerate(cards):
        x = 0.55 + i * 4.175
        badge(s, x + 1.745, 2.85, i + 1, D2)
        box(s, f'c{i}', x, 3.3, 3.9, 1.55, [h, b], A1 if i != 1 else A2, 0.0, 16, body_size=12)
    box(s, 'os', 0.55, 5.15, 6.0, 1.2, ['Runs on', 'RHEL · Ubuntu · SUSE · WSL · macOS'], D2, 0.0, 14,
        body_size=13)
    box(s, 'start', 6.8, 5.15, 6.0, 1.2, ['Two front ends, one agent', 'chipagents (terminal)   ·   '
        'chipagents web (browser)'], D1, 0.0, 14, body_size=13)
    source(s, 'Source: ChipAgents docs v0.44.3 — getting-started/overview.md')
    finish(prs, s, 'intro_ca')


def slide_intro_web(prs):
    s = content_slide(prs, 'Intro: what is ChipAgents Web?')
    text(s, 0.55, 1.45, 12.25, 0.45, 'The same agent as the terminal CLI, presented in a browser-based UI. '
         'Launched from the terminal, served as a local web app.', 14, color=GRAY, name='intro')
    box(s, 'term', 0.55, 2.2, 3.4, 1.5, ['Terminal', '$ chipagents web'], D2, 0.0, 15, body_size=13,
        font=None)
    arrow(s, 4.05, 2.75, 0.6, 0.4)
    box(s, 'server', 4.75, 2.2, 3.7, 1.5, ['Local web server', '127.0.0.1 : 4242',
        'tries next ports if busy'], A1, 0.0, 15, body_size=12)
    arrow(s, 8.55, 2.75, 0.6, 0.4)
    box(s, 'browser', 9.25, 2.2, 3.55, 1.5, ['Browser', '/hub?token=…', 'signs this browser in'], A2, 0.0,
        15, body_size=12)
    arrow(s, 6.4, 3.8, 0.4, 0.45, 'down', A1)
    box(s, 'engine', 0.55, 4.35, 12.25, 1.0, ['Same agent · tools · configuration · login as chipagents  '
        '→  behavior at parity with the terminal'], D1, 0.0, 15)
    facts = ['Context = launch folder and its children', 'First run: browser OAuth login',
             'Ctrl-C in the terminal stops the server']
    for i, f in enumerate(facts):
        box(s, f'f{i}', 0.55 + i * 4.175, 5.6, 3.9, 0.8, f, A1, 0.85, 12, bold=False, text_color=D1,
            line_color=A1)
    source(s, 'Source: ChipAgents docs v0.44.3 — usage/chipagents-web.md')
    finish(prs, s, 'intro_web')


def slide_why(prs):
    s = content_slide(prs, 'Why the web version: what I need → what Web gives')
    rows = [('Run several tasks at once', 'Multiple threads, grouped by workspace — keep one running while '
             'starting another'),
            ('Review the agent\'s edits', 'Side-pane file viewer showing the run\'s diff with +/− coloring'),
            ('Read docs, plans and diagrams as intended', 'Rendered Markdown, Mermaid and D2 diagrams in chat'),
            ('Not watch a terminal during long runs', 'Browser notifications for finished turns'),
            ('Agent in WSL, me in Windows', 'Browser UI on localhost; same agent, tools, config and login')]
    for i, (need, feat) in enumerate(rows):
        y = 1.6 + i * 0.95
        badge(s, 0.55, y + 0.18, i + 1, D2)
        box(s, f'n{i}', 1.1, y, 4.9, 0.78, need, A1, 0.85, 14, text_color=D1, align=PP_ALIGN.LEFT,
            line_color=A1)
        arrow(s, 6.1, y + 0.22, 0.5, 0.34, color=A1)
        box(s, f'f{i}', 6.7, y, 6.1, 0.78, feat, A1, 0.0, 12, bold=False, align=PP_ALIGN.LEFT)
    text(s, 0.55, 6.4, 12.25, 0.4, 'Nothing is lost: slash commands, custom commands, subagents and '
         'thinking modes work in the same session.', 14, True, A1, PP_ALIGN.CENTER, name='takeaway')
    finish(prs, s, 'why')


def slide_section(prs, title):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    set_placeholder_text(s, 10, title)
    finish(prs, s, 'section')


def slide_threads(prs):
    s = content_slide(prs, 'Feature: multiple threads and workspaces')
    container(s, 0.55, 1.95, 3.7, 4.55, 'Sidebar (illustrative)', D2)
    items = [('ws', 'Desktop'), ('th', 'Thread: web launcher'), ('th', 'Thread: deck'),
             ('ws', 'migration project'), ('th', 'Thread: DC test port'),
             ('ws', 'stdf-datatool'), ('th', 'Thread: parser update')]
    y = 2.12
    for i, (kind, label) in enumerate(items):
        if kind == 'ws':
            box(s, f'i{i}', 0.72, y, 3.36, 0.46, '▸ ' + label, D2, 0.0, 12, align=PP_ALIGN.LEFT)
            y += 0.56
        else:
            box(s, f'i{i}', 1.0, y, 3.08, 0.42, label, A1, 0.7, 11, bold=False, text_color=D1,
                align=PP_ALIGN.LEFT)
            y += 0.52
    cards = [('Parallel conversations', 'Run and switch between several threads; each keeps its own history.'),
             ('Workspace = a folder on disk', 'Add workspace from the sidebar; threads work in that folder. '
              'Name workspaces and add source folders.'),
             ('Stays in sync', 'Unread markers sync across ports and tabs; deleted threads stay deleted across '
              'browsers.'),
             ('Collaborative Projects', 'Default project per workspace (Edit workspace…); project picker per '
              'thread; current project highlighted in the header.')]
    for i, (h, b) in enumerate(cards):
        y = 1.65 + i * 1.22
        badge(s, 4.6, y + 0.3, i + 1, A1)
        box(s, f'c{i}', 5.15, y, 7.65, 1.05, [h, b], A1, 0.85, 14, text_color=D1, align=PP_ALIGN.LEFT,
            line_color=A1, body_size=12)
    source(s, 'Source: usage/chipagents-web.md; usage/projects.md; CHANGELOG (v0.44.3 docs)')
    finish(prs, s, 'threads')


def slide_diff(prs):
    s = content_slide(prs, 'Feature: file and diff viewer')
    code(s, 'chat', 0.55, 1.6, 5.3, 3.35, [
        ('> widen the free-port range', CODE_FG), ('', CODE_FG),
        ('⏺ Edit(ca-freeport.sh)', DIFF_ADD), ('  ⎿ 1 file changed', CODE_FG),
        ('    ca-freeport.sh  ← click to open', RGBColor(0x6F, 0xB3, 0xF2)), ('', CODE_FG),
        ('⏺ Done. Range is now 4242–4299.', CODE_FG)], 12, title='Chat (illustrative)')
    arrow(s, 5.95, 3.08, 0.5, 0.4, color=A1)
    code(s, 'pane', 6.55, 1.6, 6.25, 3.35, [
        ('  START="${1:-4242}"', CODE_FG), ('- END="${2:-4250}"', DIFF_DEL), ('+ END="${2:-4299}"', DIFF_ADD),
        ('', CODE_FG), ('  for ((p=START; p<=END; p++)); do', CODE_FG)], 12,
        title='Side pane — ca-freeport.sh · run diff')
    feats = [('Any file', 'Open any file the agent read or changed'),
             ('Run diff', '+/− coloring; review edits without leaving the browser'),
             ('Tree / list', 'Diff viewer views; clickable file references in messages'),
             ('Markdown', 'Rendered in the pane, toggle to editable source')]
    for i, (h, b) in enumerate(feats):
        box(s, f'f{i}', 0.55 + i * 3.1, 5.2, 2.95, 1.25, [h, b], A1 if i % 2 == 0 else A2, 0.0, 14,
            body_size=11)
    source(s, 'Source: usage/chipagents-web.md; CHANGELOG (v0.44.3 docs). Chat/diff content is a mock-up.')
    finish(prs, s, 'diff')


def slide_render(prs):
    s = content_slide(prs, 'Feature: rich rendering in the browser')
    cols = [('Syntax highlighting', A1), ('Diagrams in chat', A2), ('Transcript tools', A3)]
    for i, (h, c) in enumerate(cols):
        x = 0.55 + i * 4.175
        box(s, f'h{i}', x, 1.6, 3.9, 0.6, h, c, 0.0, 16)
    code(s, 'syn', 0.55, 2.35, 3.9, 2.4, [('def free_port(lo, hi):', RGBColor(0x6F, 0xB3, 0xF2)),
         ('    for p in range(lo, hi):', RGBColor(0xC5, 0x86, 0xC0)), ('        if ok(p):', RGBColor(0xC5, 0x86, 0xC0)),
         ('            return p', DIFF_ADD), ('    return None', CODE_FG)], 12)
    bx = 4.73
    box(s, 'diag', bx, 2.35, 3.9, 2.4, '', A2, 0.85, line_color=A2)
    for j, lbl in enumerate(['Spec', 'RTL', 'Test']):
        box(s, f'd{j}', bx + 0.15 + j * 1.27, 3.2, 0.95, 0.7, lbl, A2, 0.0, 12, tracked=False)
        if j < 2:
            arrow(s, bx + 1.13 + j * 1.27, 3.41, 0.24, 0.28, color=A2)
    text(s, bx, 4.3, 3.9, 0.35, 'rendered Mermaid / D2 (mock-up)', 10, color=GRAY, align=PP_ALIGN.CENTER)
    box(s, 'tr', 8.91, 2.35, 3.9, 2.4, ['▸ Context compacted', 'click to expand the summary', '',
        'Select text → citation'], A3, 0.85, 14, text_color=D1, line_color=A3, body_size=12)
    caps = ['Tool programs and output are highlighted; scrollback in the browser',
            'Mermaid and D2 render inline; linked diagrams refresh when the source changes; fullscreen and '
            'pinch-zoom',
            'Expand “Context compacted” to read the summary; select transcript text to create citations']
    for i, c in enumerate(caps):
        text(s, 0.55 + i * 4.175, 4.9, 3.9, 1.5, c, 12, name=f'cap{i}')
    source(s, 'Source: usage/chipagents-web.md; CHANGELOG (v0.44.3 docs). Visuals are mock-ups.')
    finish(prs, s, 'render')


def slide_productivity(prs):
    s = content_slide(prs, 'Feature: everyday productivity')
    tiles = [('Terminal tab', 'Interactive shell in the thread\'s directory'),
             ('Per-thread tabs', 'Keep each thread\'s views together'),
             ('Agents & Tasks views', 'See subagents and tasks separately'),
             ('Keyboard shortcuts', 'Plus a new-thread button when the sidebar is collapsed'),
             ('Notifications', 'Browser notification when a turn finishes'),
             ('Type or paste anywhere', 'Keystrokes go straight to the composer'),
             ('9 themes', 'Each with light and dark modes'),
             ('Preferences persist', 'Theme and layout survive restarts, even on another port')]
    for i, (h, b) in enumerate(tiles):
        r, c = divmod(i, 4)
        x, y = 0.55 + c * 3.1, 1.65 + r * 2.4
        box(s, f't{i}', x, y, 2.95, 2.2, [h, b], A1 if (r + c) % 2 == 0 else A2, 0.0, 16, body_size=12)
    source(s, 'Source: usage/chipagents-web.md; CHANGELOG (v0.44.3 docs)')
    finish(prs, s, 'productivity')


def slide_control(prs):
    s = content_slide(prs, 'Feature: same engine and controls as the CLI, plus sharing')
    container(s, 0.55, 1.95, 6.0, 4.5, 'Same engine as the CLI', A1)
    left = [('Full CLI feature set', 'Slash commands, custom commands, subagents, thinking modes'),
            ('/provider', 'Choose a model and save defaults'),
            ('Auto review mode', 'Default for new web sessions; pending approvals shown in the UI'),
            ('Same sandbox and hooks', 'Asks before reaching a domain outside the allow list; '
             'workflow hooks run as in the CLI')]
    for i, (h, b) in enumerate(left):
        box(s, f'l{i}', 0.75, 2.12 + i * 1.07, 5.6, 0.95, [h, b], A1, 0.0, 13, align=PP_ALIGN.LEFT,
            body_size=11)
    container(s, 6.8, 1.95, 6.0, 4.5, 'Collaboration and modes', A2)
    right = [('Share button and dialogs', 'Share threads, including with user groups'),
             ('/share', 'Copies the thread\'s dashboard link'),
             ('Project picker', 'Assign a Collaborative Project per thread'),
             ('Mode selector', 'Switch between chipagents web and chipagents formal')]
    for i, (h, b) in enumerate(right):
        box(s, f'r{i}', 7.0, 2.12 + i * 1.07, 5.6, 0.95, [h, b], A2, 0.0, 13, align=PP_ALIGN.LEFT,
            body_size=11)
    source(s, 'Source: usage/chipagents-web.md; usage/projects.md; configuration/sandboxing.md; CHANGELOG '
           '(v0.44.3 docs)')
    finish(prs, s, 'control')


def slide_security(prs):
    s = content_slide(prs, 'Security model: three gates before a session')
    box(s, 'req', 0.55, 2.0, 1.9, 1.3, ['Request', 'from a browser'], D2, 0.0, 14, body_size=11)
    gates = [('Loopback only', '127.0.0.1', 'Not reachable from other machines; use SSH -L for remote hosts'),
             ('Access token', '/hub?token=…', 'Valid for the life of the server; treat the link like a password'),
             ('Same Unix user', 'Linux', 'Connecting process must run as the same user; else 403')]
    for i, (h, sub, desc) in enumerate(gates):
        x = 3.0 + i * 2.85
        arrow(s, x - 0.48, 2.45, 0.4, 0.4)
        box(s, f'g{i}', x, 2.0, 2.4, 1.3, [h, sub], A1, 0.0, 14, body_size=11)
        text(s, x, 3.45, 2.4, 1.2, desc, 11, align=PP_ALIGN.CENTER, name=f'gd{i}')
    arrow(s, 11.07, 2.45, 0.4, 0.4, color=GREEN)
    box(s, 'ses', 11.55, 2.0, 1.25, 1.3, ['Session', 'granted'], GREEN, size=13, body_size=11)
    box(s, 'multi', 0.55, 4.7, 12.25, 0.65, 'Several people can run their own chipagents web on one host '
        'without reaching each other\'s sessions.', A1, 0.85, 13, bold=False, text_color=D1, line_color=A1)
    box(s, 'warn', 0.55, 5.55, 12.25, 0.75, 'With CHIPAGENTS_DISABLE_UID_CHECK=true the token and loopback '
        'checks remain, but anyone who obtains the link can use the session — keep it private.', AMBER,
        size=12, bold=False)
    source(s, 'Source: usage/chipagents-web.md (v0.44.3 docs)')
    finish(prs, s, 'security')


def slide_setup(prs):
    s = content_slide(prs, 'My setup: one double-click from Windows')
    steps = [('chipagents-web.bat', 'Windows double-click'), ('ca-freeport.sh', 'first free port\n4242–4299'),
             ('WSL bash', 'environment setup'), ('chipagents web', 'HUB_PORT=<port>'),
             ('Browser', 'localhost:<port>/hub')]
    for i, (h, b) in enumerate(steps):
        x = 0.55 + i * 2.5
        badge(s, x + 0.9, 1.55, i + 1, D2)
        box(s, f's{i}', x, 2.0, 2.15, 1.3, [h] + b.split('\n'), A1 if i % 2 == 0 else A2, 0.0, 13,
            body_size=11)
        if i < len(steps) - 1:
            arrow(s, x + 2.19, 2.5, 0.27, 0.3)
    code(s, 'env', 0.55, 3.6, 7.3, 2.0, [
        ('cd /mnt/c/Users/jcortez3/Desktop', CODE_FG),
        ('export NODE_EXTRA_CA_CERTS=$HOME/certs/adi-ca-bundle.pem', CODE_FG),
        ('export CHIPAGENTS_DISABLE_UID_CHECK=true', RGBColor(0xF2, 0xC2, 0x6D)),
        ('HUB_PORT=$PORT ~/.local/bin/chipagents web', DIFF_ADD)], 11, title='What the launcher runs in WSL')
    box(s, 'why', 8.1, 3.6, 4.7, 2.0, ['Why the UID check is off', 'The Windows browser reaching WSL got '
        '403 "Connections from other users are not allowed". Disabling the check is the documented fix.'],
        AMBER, 0.0, 13, bold=True, align=PP_ALIGN.LEFT, body_size=11)
    box(s, 'tip', 0.55, 5.8, 12.25, 0.65, 'Sign in once with the access link printed in the console; '
        'the tab then stays signed in until the server stops.', A1, 0.85, 12, bold=False, text_color=D1,
        line_color=A1)
    finish(prs, s, 'setup')


def slide_limits(prs):
    s = content_slide(prs, 'Notes and limitations (v0.44.3)')
    items = [('No image paste in the chat box', 'Save the screenshot and reference its path in the prompt'),
             ('New link per server start', 'The token is valid only while the server runs'),
             ('UID check disabled in my WSL setup', 'Keep the access link private'),
             ('Waveforms', 'Viewed with the separate chipagents wave-viewer command'),
             ('Closing the console stops the server', 'Ctrl-C or closing the window ends all web sessions')]
    for i, (h, b) in enumerate(items):
        y = 1.6 + i * 0.95
        badge(s, 0.55, y + 0.18, i + 1, AMBER)
        box(s, f'i{i}', 1.1, y, 11.7, 0.78, [h, b], A1, 0.88, 14, text_color=D1, align=PP_ALIGN.LEFT,
            line_color=D2, body_size=12)
    text(s, 0.55, 6.4, 12.25, 0.4, 'Bottom line: same agent as the CLI, with a better workspace for '
         'multi-task work.', 15, True, A1, PP_ALIGN.CENTER, name='takeaway')
    finish(prs, s, 'limits')


def slide_end(prs):
    s = prs.slides.add_slide(prs.slide_layouts[21])
    remove_unused_placeholders(s)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--output-dir', default=os.path.dirname(os.path.abspath(__file__)))
    args = ap.parse_args()
    tpl = TEMPLATE_PATH
    if tpl.endswith('.potx'):
        tpl = convert_potx_to_pptx(tpl)
    prs = Presentation(tpl)
    clear_slides(prs)
    slide_cover(prs)
    slide_intro_ca(prs)
    slide_intro_web(prs)
    slide_why(prs)
    slide_section(prs, 'Features')
    slide_threads(prs)
    slide_diff(prs)
    slide_render(prs)
    slide_productivity(prs)
    slide_control(prs)
    slide_security(prs)
    slide_section(prs, 'How I run it')
    slide_setup(prs)
    slide_limits(prs)
    slide_end(prs)
    out = os.path.join(args.output_dir, 'Why_ChipAgents_Web.pptx')
    prs.save(out)
    print(f'Saved {out}')
    verify_pptx(out, expected_slides=15)


if __name__ == '__main__':
    main()
