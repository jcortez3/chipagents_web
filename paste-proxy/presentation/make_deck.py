import argparse
import os
import sys

SKILL_SCRIPTS = os.path.expanduser('~/.config/chipagents/config2/skills/pptx-generator/scripts')
sys.path.insert(0, SKILL_SCRIPTS)

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from pptx_template import (TEMPLATE_PATH, convert_potx_to_pptx, clear_slides, get_placeholder,
                           remove_unused_placeholders, apply_footer, set_placeholder_text, verify_pptx)

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
RED = RGBColor(0xB3, 0x26, 0x1E)
AMBER = RGBColor(0xB2, 0x6A, 0x00)
GRAY = RGBColor(0x60, 0x60, 0x60)
MONO = 'Consolas'
A1, A2, A3, A4 = (MSO_THEME_COLOR.ACCENT_1, MSO_THEME_COLOR.ACCENT_2,
                  MSO_THEME_COLOR.ACCENT_3, MSO_THEME_COLOR.ACCENT_4)
BOXES = []


def reset_boxes():
    BOXES.clear()


def track(name, x, y, w, h):
    BOXES.append((name, x, y, w, h))


def check_layout(slide_name):
    for i, (n1, x1, y1, w1, h1) in enumerate(BOXES):
        if x1 < 0.3 or y1 < 1.3 or x1 + w1 > 13.05 or y1 + h1 > 6.95:
            print(f'  WARN [{slide_name}] {n1} outside safe area')
        for n2, x2, y2, w2, h2 in BOXES[i + 1:]:
            if not (x1 + w1 <= x2 or x2 + w2 <= x1 or y1 + h1 <= y2 or y2 + h2 <= y1):
                print(f'  WARN [{slide_name}] {n1} overlaps {n2}')
    reset_boxes()


def _style_text(tf, lines, size, bold, color, align, font=None):
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.font.size = Pt(size)
        p.font.bold = bold if i == 0 else False
        if font:
            p.font.name = font
        if isinstance(color, RGBColor):
            p.font.color.rgb = color
        else:
            p.font.color.theme_color = color


def box(slide, name, x, y, w, h, lines, fill, brightness=0.0, size=13, bold=True,
        text_color=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, font=None, align=PP_ALIGN.CENTER,
        line_color=None, tracked=True):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    if isinstance(fill, RGBColor):
        s.fill.fore_color.rgb = fill
    else:
        s.fill.fore_color.theme_color = fill
        s.fill.fore_color.brightness = brightness
    if line_color is None:
        s.line.fill.background()
    else:
        s.line.color.theme_color = line_color
        s.line.width = Pt(1.5)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.anchor = MSO_ANCHOR.MIDDLE
    _style_text(tf, lines if isinstance(lines, list) else [lines], size, bold, text_color, align, font)
    if tracked:
        track(name, x, y, w, h)
    return s


def container(slide, name, x, y, w, h, label, color):
    lbl = slide.shapes.add_textbox(Inches(x), Inches(y - 0.36), Inches(w), Inches(0.32))
    _style_text(lbl.text_frame, [label], 13, True, color, PP_ALIGN.LEFT)
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.04
    s.fill.solid()
    s.fill.fore_color.theme_color = color
    s.fill.fore_color.brightness = 0.85
    s.line.color.theme_color = color
    s.line.width = Pt(1.25)
    return s


def text(slide, x, y, w, h, lines, size=11, bold=False, color=MSO_THEME_COLOR.DARK_1,
         align=PP_ALIGN.LEFT, font=None, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.text_frame.margin_left = tb.text_frame.margin_right = Inches(0.02)
    _style_text(tb.text_frame, lines if isinstance(lines, list) else [lines], size, bold, color, align, font)
    if name:
        track(name, x, y, w, h)
    return tb


def arrow(slide, x, y, w, h, direction='right', color=MSO_THEME_COLOR.DARK_2):
    shp = {'right': MSO_SHAPE.RIGHT_ARROW, 'left': MSO_SHAPE.LEFT_ARROW,
           'down': MSO_SHAPE.DOWN_ARROW, 'up': MSO_SHAPE.UP_ARROW}[direction]
    a = slide.shapes.add_shape(shp, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.theme_color = color
    a.line.fill.background()
    return a


def badge(slide, x, y, n, color=A1, d=0.42):
    b = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    b.fill.solid()
    b.fill.fore_color.theme_color = color
    b.line.color.rgb = WHITE
    b.line.width = Pt(1.5)
    tf = b.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.anchor = MSO_ANCHOR.MIDDLE
    _style_text(tf, [str(n)], 13, True, WHITE, PP_ALIGN.CENTER)
    return b


def content_slide(prs, title, layout=2):
    slide = prs.slides.add_slide(prs.slide_layouts[layout])
    set_placeholder_text(slide, 0, title, font_size=26, bold=True)
    for idx in (10, 11):
        ph = get_placeholder(slide, idx)
        if ph is not None:
            ph._element.getparent().remove(ph._element)
    return slide


def finish(prs, slide, name):
    apply_footer(slide, prs)
    remove_unused_placeholders(slide)
    check_layout(name)


def table(slide, name, x, y, w, col_w, headers, rows, size=11, row_h=0.36, status_col=None):
    t = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w),
                               Inches(row_h * (len(rows) + 1))).table
    for i, cw in enumerate(col_w):
        t.columns[i].width = Inches(cw)
    for ci, h in enumerate(headers):
        c = t.cell(0, ci)
        c.text = h
        p = c.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(size + 1)
        p.font.color.rgb = WHITE
        c.fill.solid()
        c.fill.fore_color.theme_color = A1
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.cell(ri + 1, ci)
            c.text = val
            p = c.text_frame.paragraphs[0]
            p.font.size = Pt(size)
            p.font.color.theme_color = MSO_THEME_COLOR.DARK_1
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            if ri % 2 == 0:
                c.fill.fore_color.theme_color = A1
                c.fill.fore_color.brightness = 0.85
            else:
                c.fill.fore_color.rgb = WHITE
            if status_col is not None and ci == status_col:
                p.font.bold = True
                p.alignment = PP_ALIGN.CENTER
                p.font.color.rgb = {'PASS': GREEN, 'PENDING': AMBER}.get(val.split()[0], RED)
    track(name, x, y, w, row_h * (len(rows) + 1))
    return t


def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    set_placeholder_text(slide, 10, 'Image Paste for ChipAgents Web')
    set_placeholder_text(slide, 14, 'Local proxy workaround · WSL2')
    set_placeholder_text(slide, 15, 'ChipAgents Web v0.44.3')
    finish(prs, slide, 'cover')


def slide_problem(prs):
    s = content_slide(prs, 'Problem: screenshots take five steps to reach the agent')
    text(s, 0.55, 1.45, 12.2, 0.4, 'Today (v0.44.3): the chat box accepts text only. Images must be saved '
         'and referenced by path.', 14, color=GRAY, name='intro')
    container(s, 'before', 0.55, 2.35, 12.25, 1.55, 'Before — manual', MSO_THEME_COLOR.DARK_2)
    steps = ['Win+Shift+S\nscreenshot', 'Open Paint /\nSave As…', 'Pick folder\n+ file name',
             'Type /mnt/c/…\npath in chat', 'Agent reads\nthe file']
    for i, st in enumerate(steps):
        x = 0.8 + i * 2.42
        box(s, f'b{i}', x, 2.6, 2.0, 1.05, st.split('\n'), MSO_THEME_COLOR.DARK_2, 0.0, 12, bold=False,
            tracked=False)
        if i < len(steps) - 1:
            arrow(s, x + 2.04, 2.98, 0.34, 0.3)
    container(s, 'after', 0.55, 4.55, 12.25, 1.55, 'After — with the paste proxy', A1)
    box(s, 'a0', 0.8, 4.8, 2.0, 1.05, ['Win+Shift+S', 'screenshot'], A1, 0.0, 12, bold=False, tracked=False)
    arrow(s, 2.84, 5.18, 0.34, 0.3, color=A1)
    box(s, 'a1', 3.22, 4.8, 2.0, 1.05, ['Ctrl+V', 'in chat box'], A1, 0.0, 12, bold=False, tracked=False)
    arrow(s, 5.26, 5.18, 0.34, 0.3, color=A1)
    box(s, 'a2', 5.64, 4.8, 4.42, 1.05, ['Path inserted automatically',
        '[pasted image: /mnt/c/…/screens/paste_<ts>.png]'], A1, 0.0, 12, bold=False, tracked=False)
    arrow(s, 10.1, 5.18, 0.34, 0.3, color=A1)
    box(s, 'a3', 10.48, 4.8, 2.0, 1.05, ['Agent reads', 'the file'], A1, 0.0, 12, bold=False, tracked=False)
    text(s, 0.55, 6.3, 12.2, 0.4, '5 manual steps  →  2 keystrokes', 16, True, A1, PP_ALIGN.CENTER,
         name='takeaway')
    finish(prs, s, 'problem')


def slide_architecture(prs):
    s = content_slide(prs, 'Architecture: a pass-through proxy in front of chipagents web')
    container(s, 'win', 0.5, 1.95, 2.95, 4.25, 'Windows', MSO_THEME_COLOR.DARK_2)
    box(s, 'browser', 0.7, 2.25, 2.55, 1.35, ['Browser', 'Edge / Chrome', 'localhost:4300/hub?token=…'],
        MSO_THEME_COLOR.DARK_2, 0.0, 12)
    box(s, 'inject', 0.7, 4.0, 2.55, 1.35, ['paste-inject.js', 'paste handler injected', 'into the hub page'],
        A2, 0.0, 12)
    text(s, 0.7, 5.45, 2.55, 0.6, ['Clipboard image →', 'upload + insert path'], 10, color=GRAY,
         align=PP_ALIGN.CENTER, name='inject_note')
    arrow(s, 3.55, 2.72, 0.75, 0.4)
    text(s, 3.35, 3.15, 1.2, 0.75, ['HTTP', 'SSE', 'WebSocket'], 10, True, MSO_THEME_COLOR.DARK_2,
         PP_ALIGN.CENTER, name='proto')
    container(s, 'wsl', 4.4, 1.95, 8.45, 4.25, 'WSL2 · Unix user jcortez3', A1)
    box(s, 'proxy', 4.65, 2.25, 3.2, 1.35, ['paste-proxy.js', 'Node 18 · 127.0.0.1:4300',
        'injects script · forwards traffic'], A1, 0.0, 12)
    arrow(s, 7.95, 2.72, 0.75, 0.4, color=A1)
    text(s, 7.88, 3.15, 0.9, 0.5, ['token', 'forwarded'], 10, True, A1, PP_ALIGN.CENTER, name='fwd')
    box(s, 'ca', 8.8, 2.25, 3.8, 1.35, ['chipagents web v0.44.3', '127.0.0.1:4242',
        'child process of the proxy'], MSO_THEME_COLOR.DARK_1, 0.0, 12)
    arrow(s, 6.05, 3.68, 0.4, 0.6, 'down', A2)
    text(s, 6.5, 3.75, 2.2, 0.45, 'POST /__paste/upload', 10, True, A2, font=MONO, name='post')
    box(s, 'screens', 4.65, 4.35, 3.2, 1.35, ['Desktop\\screens\\', 'paste_<timestamp>.png'], A2, 0.0, 12)
    arrow(s, 7.95, 4.83, 0.75, 0.4, color=A3)
    text(s, 7.88, 5.25, 0.9, 0.3, 'View tool', 10, True, A3, PP_ALIGN.CENTER, name='reads')
    box(s, 'agent', 8.8, 4.35, 3.8, 1.35, ['Agent', 'reads the image from', '/mnt/c/…/screens/…'], A3, 0.0, 12)
    text(s, 0.5, 6.32, 12.35, 0.4, 'Only two routes are new (/__paste/*). Everything else passes through '
         'unchanged.', 14, True, A1, PP_ALIGN.CENTER, name='takeaway')
    finish(prs, s, 'architecture')


def slide_flow(prs):
    s = content_slide(prs, 'Paste flow: from Ctrl+V to a file path in the prompt')
    steps = [('Ctrl+V', 'in chat box'), ('Intercept', 'capture-phase\npaste listener'),
             ('Check', 'image file &\nno plain text'), ('Upload', 'POST bytes\n/__paste/upload'),
             ('Save', 'screens\\paste_\n<ts>.<ext>'), ('Insert', 'path at cursor\n+ green toast'),
             ('Send', 'agent reads\nimage')]
    w, gap, x0 = 1.5, 0.29, 0.55
    for i, (head, body) in enumerate(steps):
        x = x0 + i * (w + gap)
        box(s, f's{i}', x, 1.95, w, 1.3, [head] + body.split('\n'), A1 if i % 2 == 0 else A2, 0.0, 11,
            tracked=False)
        badge(s, x + w / 2 - 0.21, 1.5, i + 1, MSO_THEME_COLOR.DARK_2)
        if i < len(steps) - 1:
            arrow(s, x + w + 0.035, 2.47, 0.22, 0.26)
    track('flow', 0.55, 1.5, 7 * w + 6 * gap, 1.75)
    text(s, 0.55, 3.55, 6.0, 0.35, 'Interception rule', 14, True, A1, name='rule_h')
    box(s, 'rule', 0.55, 3.95, 6.0, 1.6, ['Handle the paste only when ALL are true:',
        '• focus is a textarea, text input or contenteditable',
        '• clipboard holds a png / jpg / gif / webp file',
        '• clipboard has no plain text (Excel/Word copies pass through)'],
        A1, 0.85, 12, text_color=MSO_THEME_COLOR.DARK_1, align=PP_ALIGN.LEFT, line_color=A1)
    text(s, 6.85, 3.55, 5.95, 0.35, 'What you see in the chat box', 14, True, A1, name='ui_h')
    box(s, 'ui', 6.85, 3.95, 5.95, 1.0, ['[pasted image: /mnt/c/Users/jcortez3/Desktop/screens/'
        'paste_20261029_172044_855.png] why is this test failing?'], MSO_THEME_COLOR.LIGHT_1, 0.0, 11,
        bold=False, text_color=MSO_THEME_COLOR.DARK_1, font=MONO, align=PP_ALIGN.LEFT,
        line_color=MSO_THEME_COLOR.DARK_2)
    box(s, 'toast', 9.15, 5.1, 3.65, 0.45, 'Image saved: paste_20261029_172044_855.png', GREEN, size=10,
        bold=False)
    text(s, 0.55, 5.85, 12.25, 0.75, ['Insertion uses document.execCommand("insertText"), so the React '
         'input state and Ctrl+Z undo stay consistent.', 'Fallback: native value setter + input event.'],
         11, color=GRAY, name='note')
    finish(prs, s, 'flow')


def slide_uid(prs):
    s = content_slide(prs, 'Side effect: the Windows "other users" 403 goes away')
    rows = [('Before — direct to chipagents web', 2.0,
             [('Windows browser', MSO_THEME_COLOR.DARK_2), ('WSL localhost\nforwarding', MSO_THEME_COLOR.DARK_2),
              ('chipagents web\n:4242', MSO_THEME_COLOR.DARK_1)],
             ('403  "Connections from other\nusers are not allowed"', RED)),
            ('After — through the paste proxy', 4.15,
             [('Windows browser', MSO_THEME_COLOR.DARK_2), ('WSL localhost\nforwarding', MSO_THEME_COLOR.DARK_2),
              ('paste-proxy.js\n(runs as jcortez3)', A1), ('chipagents web\n:4242', MSO_THEME_COLOR.DARK_1)],
             ('200  same Unix\nuser → accepted', GREEN))]
    for label, y, blocks, (res, rc) in rows:
        text(s, 0.55, y - 0.42, 8.0, 0.35, label, 14, True, A1, name=f'lbl{y}')
        x = 0.55
        bw = 2.15
        for i, (t, c) in enumerate(blocks):
            box(s, f'{y}_{i}', x, y, bw, 1.15, t.split('\n'), c, 0.0, 12)
            arrow(s, x + bw + 0.06, y + 0.4, 0.38, 0.35, color=MSO_THEME_COLOR.DARK_2)
            x += bw + 0.5
        box(s, f'{y}_res', x, y, 12.8 - x, 1.15, res.split('\n'), rc, size=12)
    text(s, 0.55, 5.55, 12.25, 1.2, [
        'Observed: a Windows request straight to chipagents web returns 403. The same request through the '
        'proxy returns 200 on /hub/api/workspaces.',
        'Why: the server checks that the connecting process runs as the same Unix user. The proxy runs '
        'inside WSL as jcortez3. The exact cause of the direct-path failure is inferred, not documented.',
        'Result: CHIPAGENTS_DISABLE_UID_CHECK is not needed with this launcher. Token auth is still enforced.'],
        11, color=MSO_THEME_COLOR.DARK_1, name='notes')
    finish(prs, s, 'uid')


def slide_security(prs):
    s = content_slide(prs, 'Security controls')
    rows = [
        ('Network exposure', 'Proxy binds 127.0.0.1 only; upstream is 127.0.0.1:4242', 'No LAN/remote access'),
        ('Authentication', 'Hub token (x-access-token / ?token=) forwarded untouched',
         'Missing token → 401 from chipagents'),
        ('Cross-site upload', 'Upload requires X-ChipAgents-Paste: 1 header',
         'Forces CORS preflight, which the proxy never approves'),
        ('Origin check', 'Origin must be http://localhost|127.0.0.1:<proxy port>', 'Other origins → 403'),
        ('Content type', 'png, jpg, gif, webp only', 'Others → 415'),
        ('Size limit', '25 MB per image', 'Larger → 413'),
        ('Data egress', 'Files stay on the local disk (Desktop\\screens)', 'No external network calls'),
        ('UID check', 'Left enabled; proxy connects as the same Unix user', 'No env override needed'),
    ]
    table(s, 'tbl', 0.55, 1.55, 12.25, [2.4, 5.6, 4.25], ['Control', 'Implementation', 'Effect'], rows,
          size=12, row_h=0.5)
    text(s, 0.55, 6.25, 12.25, 0.45, 'Not covered: other local users who obtain the hub token. Keep the '
         'printed link private, as with plain chipagents web.', 12, color=GRAY, name='note')
    finish(prs, s, 'security')


def slide_verification(prs):
    s = content_slide(prs, 'Verification results')
    rows = [
        ('Hub HTML via proxy', 'GET /hub (Accept: text/html)', 'inject.js tag present', 'PASS'),
        ('Inject script served', 'GET /__paste/inject.js', 'HTTP 200', 'PASS'),
        ('Large asset pass-through', 'GET /hub/main.js', 'HTTP 200, 19,559,835 B', 'PASS'),
        ('Upload without header', 'POST /__paste/upload', 'HTTP 403 forbidden', 'PASS'),
        ('Upload, foreign origin', 'Origin: http://evil.com', 'HTTP 403 forbidden', 'PASS'),
        ('Upload, valid', 'PNG + header + loopback origin', 'File saved, path returned', 'PASS'),
        ('Auth preserved', 'GET /api/version, no token', 'HTTP 401', 'PASS'),
        ('Windows → proxy (auth)', 'curl.exe /hub/api/workspaces + token', 'HTTP 200', 'PASS'),
        ('Browser Ctrl+V end-to-end', 'Edge/Chrome, Win+Shift+S → Ctrl+V', 'Path inserted, toast shown',
         'PENDING'),
    ]
    table(s, 'tbl', 0.55, 1.5, 12.25, [3.1, 3.9, 3.65, 1.6], ['Check', 'Stimulus', 'Expected / observed',
          'Result'], rows, size=11, row_h=0.47, status_col=3)
    text(s, 0.55, 6.35, 12.25, 0.4, 'All checks except the last were run with curl from WSL and from '
         'Windows. The browser paste test is the next manual step.', 12, color=GRAY, name='note')
    finish(prs, s, 'verification')


def slide_usage(prs):
    s = content_slide(prs, 'Files and usage')
    text(s, 0.55, 1.5, 6.0, 0.35, 'Files  (Desktop\\chiagents web\\)', 14, True, A1, name='fh')
    files = [('chipagents-web-paste.bat', 'One-click launcher (WSL, CA bundle, ports)'),
             ('paste-proxy\\paste-proxy.js', 'Proxy + child chipagents web, no npm deps'),
             ('paste-proxy\\paste-inject.js', 'Browser-side paste handler'),
             ('paste-proxy\\README.md', 'How it works, env vars, limits')]
    for i, (f, d) in enumerate(files):
        y = 1.95 + i * 0.82
        box(s, f'f{i}', 0.55, y, 6.0, 0.7, [f, d], A1, 0.85, 11, text_color=MSO_THEME_COLOR.DARK_1,
            align=PP_ALIGN.LEFT, line_color=A1)
    text(s, 6.95, 1.5, 5.85, 0.35, 'Run', 14, True, A1, name='rh')
    for i, st in enumerate(['Double-click chipagents-web-paste.bat',
                            'Browser opens the tokenized link on port 4300',
                            'Win+Shift+S, click the chat box, Ctrl+V']):
        y = 1.95 + i * 0.6
        badge(s, 6.95, y + 0.04, i + 1, A2, 0.38)
        text(s, 7.45, y, 5.35, 0.45, st, 12, name=f'run{i}')
    text(s, 6.95, 3.85, 5.85, 0.35, 'Environment variables', 14, True, A1, name='eh')
    table(s, 'env', 6.95, 4.25, 5.85, [2.35, 3.5], ['Variable', 'Default'],
          [('PASTE_DIR', '/mnt/c/…/Desktop/screens'), ('PASTE_PROXY_PORT', '4300 (first tried)'),
           ('HUB_PORT', '4242 (chipagents start port)'), ('PASTE_NO_BROWSER', 'unset (1 = no auto-open)')],
          size=10, row_h=0.42)
    finish(prs, s, 'usage')


def slide_limits(prs):
    s = content_slide(prs, 'Limitations and next steps')
    text(s, 0.55, 1.5, 6.0, 0.35, 'Limitations', 14, True, RED, name='lh')
    lim = ['Unofficial: a ChipAgents update can change the hub page and break the injection',
           'Use the proxy link (port 4300); the direct link has no paste and hits the UID 403',
           'Pasted images are not deleted automatically; clean Desktop\\screens periodically',
           'Clipboard paste only; no drag-and-drop yet',
           'Agent sees a file path, not an inline thumbnail in the chat history']
    for i, t in enumerate(lim):
        box(s, f'l{i}', 0.55, 1.95 + i * 0.86, 6.0, 0.74, t, A1, 0.9, 11, bold=False,
            text_color=MSO_THEME_COLOR.DARK_1, align=PP_ALIGN.LEFT, line_color=MSO_THEME_COLOR.DARK_2)
    text(s, 6.95, 1.5, 5.85, 0.35, 'Next steps', 14, True, GREEN, name='nh')
    nxt = ['Run the browser Ctrl+V end-to-end test (only pending check)',
           'Request native image paste via /bug-report or the ChipAgents admin',
           'Optional: auto-delete screens older than N days',
           'Optional: drag-and-drop images onto the chat box',
           'Retire the proxy once ChipAgents supports image paste']
    for i, t in enumerate(nxt):
        badge(s, 6.95, 2.11 + i * 0.86, i + 1, A2, 0.42)
        box(s, f'n{i}', 7.5, 1.95 + i * 0.86, 5.3, 0.74, t, A2, 0.85, 11, bold=False,
            text_color=MSO_THEME_COLOR.DARK_1, align=PP_ALIGN.LEFT, line_color=A2)
    finish(prs, s, 'limits')


def slide_end(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[21])
    remove_unused_placeholders(slide)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--output-dir', default=os.path.dirname(os.path.abspath(__file__)))
    args = ap.parse_args()
    tpl = TEMPLATE_PATH
    if tpl.endswith('.potx'):
        tpl = convert_potx_to_pptx(tpl)
    prs = Presentation(tpl)
    clear_slides(prs)
    for fn in (slide_cover, slide_problem, slide_architecture, slide_flow, slide_uid, slide_security,
               slide_verification, slide_usage, slide_limits, slide_end):
        fn(prs)
    out = os.path.join(args.output_dir, 'ChipAgents_Web_Image_Paste.pptx')
    prs.save(out)
    print(f'Saved {out}')
    verify_pptx(out, expected_slides=10)


if __name__ == '__main__':
    main()
